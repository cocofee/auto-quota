"""
Unified parser for priced bill documents.

The parser focuses on extracting bill-level unit prices from historical priced
files so the unified price reference layer can answer composite price queries.
"""

from __future__ import annotations

from dataclasses import asdict, dataclass, field
from pathlib import Path
import re
import xml.etree.ElementTree as ET

import openpyxl

try:
    from lxml import etree as LET
except ImportError:  # optional dependency for malformed XML recovery
    LET = None

from src.bill_reader import BillReader, _is_material_code, _is_quota_code


PRICE_ATTR_KEYS = (
    "综合单价",
    "单价",
    "含税综合单价",
    "不含税综合单价",
    "综合单价(元)",
    "ZHDJ",
    "DJ",
    "HSZHDJ",
    "BHSZHDJ",
)

ZJ_BILL_TAGS = {
    "\u5206\u90e8\u5206\u9879\u5de5\u7a0b\u91cf\u6e05\u5355\u8868\u8bb0\u5f55",
    "\u6280\u672f\u63aa\u65bd\u9879\u76ee\u6e05\u5355\u8868\u8bb0\u5f55",
}

ZJ_ANALYSIS_TAGS = {
    "\u5206\u90e8\u5206\u9879\u7efc\u5408\u5355\u4ef7\u5206\u6790\u8868",
    "\u6280\u672f\u63aa\u65bd\u6e05\u5355\u7efc\u5408\u5355\u4ef7\u5206\u6790\u8868",
}

ZJ_PROJECT_NAME_TAG = "\u5efa\u8bbe\u9879\u76ee\u4fe1\u606f\u8868"
ZJ_UNIT_TABLE_TAG = "\u5355\u4f4d\u5de5\u7a0b\u5217\u8868"
ZJ_SPECIALTY_TAG = "\u4e13\u4e1a\u5de5\u7a0b\u5217\u8868"

ZJ_ATTR_KEYS = {
    "project_name": ("\u9879\u76ee\u540d\u79f0",),
    "unit_name": ("\u5355\u4f4d\u5de5\u7a0b\u540d\u79f0",),
    "specialty_name": ("\u4e13\u4e1a\u5de5\u7a0b\u540d\u79f0",),
    "specialty_type": ("\u4e13\u4e1a\u7c7b\u578b",),
    "section_name": ("\u540d\u79f0",),
    "boq_code": ("\u9879\u76ee\u7f16\u7801",),
    "boq_name": ("\u9879\u76ee\u540d\u79f0",),
    "feature_text": ("\u9879\u76ee\u7279\u5f81",),
    "unit": ("\u8ba1\u91cf\u5355\u4f4d",),
    "quantity": ("\u5de5\u7a0b\u91cf", "\u5de5\u7a0b\u6570\u91cf"),
    "composite_unit_price": ("\u7efc\u5408\u5355\u4ef7",),
    "quota_code": ("\u7f16\u7801",),
    "quota_name": ("\u540d\u79f0",),
    "quota_unit": ("\u5355\u4f4d",),
    "quota_quantity": ("\u6570\u91cf",),
}

COST_ATTR_MAP = {
    "labor_cost": ("人工费", "人工合价", "RGF"),
    "material_cost": ("材料费", "材料合价", "CLF"),
    "machine_cost": ("机械费", "机械合价", "JXF"),
    "management_fee": ("企业管理费", "管理费", "GLF"),
    "profit": ("利润", "LR"),
    "measure_fee": ("措施费", "CSF"),
    "other_fee": ("其他费", "风险费用", "QTF"),
    "tax": ("税金", "SJ"),
}


def _clean_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _coerce_float(value) -> float | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    text = text.replace(",", "")
    text = text.replace("，", "")
    try:
        return float(text)
    except (TypeError, ValueError):
        return None


def _find_attr(attrs: dict[str, str], *keys: str) -> str:
    for key in keys:
        value = attrs.get(key)
        if value not in (None, ""):
            return str(value)
    return ""


def _find_numeric_attr(attrs: dict[str, str], *keys: str) -> float | None:
    for key in keys:
        value = attrs.get(key)
        number = _coerce_float(value)
        if number is not None:
            return number
    return None


def _normalize_name(name: str) -> str:
    return re.sub(r"\s+", " ", _clean_text(name))


def _extract_search_text(item_name: str, feature_text: str, quota_codes: list[str], quota_names: list[str]) -> str:
    parts = [item_name, feature_text, " ".join(quota_codes), " ".join(quota_names)]
    return " ".join(part for part in parts if part).strip()


@dataclass
class ParsedQuota:
    code: str = ""
    name: str = ""
    unit: str = ""
    quantity: float | None = None


@dataclass
class ParsedMaterial:
    code: str = ""
    name: str = ""
    unit: str = ""
    price: float | None = None


@dataclass
class ParsedPricedBillItem:
    boq_code: str = ""
    boq_name_raw: str = ""
    boq_name_normalized: str = ""
    feature_text: str = ""
    unit: str = ""
    quantity: float | None = None
    composite_unit_price: float | None = None
    quota_code: str = ""
    quota_name: str = ""
    specialty: str = ""
    system_name: str = ""
    subsystem_name: str = ""
    source_sheet: str = ""
    source_row_no: int | None = None
    search_text: str = ""
    bill_text: str = ""
    materials: list[ParsedMaterial] = field(default_factory=list)
    quotas: list[ParsedQuota] = field(default_factory=list)
    labor_cost: float | None = None
    material_cost: float | None = None
    machine_cost: float | None = None
    management_fee: float | None = None
    profit: float | None = None
    measure_fee: float | None = None
    other_fee: float | None = None
    tax: float | None = None
    remarks: str = ""
    tags: list[str] = field(default_factory=list)

    def to_record(self) -> dict:
        quota_code = self.quota_code
        quota_name = self.quota_name
        if not quota_code and self.quotas:
            quota_code = ",".join(q.code for q in self.quotas if q.code)
        if not quota_name and self.quotas:
            quota_name = " + ".join(q.name for q in self.quotas if q.name)
        record = {
            "boq_code": self.boq_code,
            "boq_name_raw": self.boq_name_raw,
            "boq_name_normalized": self.boq_name_normalized or self.boq_name_raw,
            "feature_text": self.feature_text,
            "unit": self.unit,
            "quantity": self.quantity,
            "composite_unit_price": self.composite_unit_price,
            "quota_code": quota_code,
            "quota_name": quota_name,
            "specialty": self.specialty,
            "system_name": self.system_name,
            "subsystem_name": self.subsystem_name,
            "source_sheet": self.source_sheet,
            "source_row_no": self.source_row_no,
            "search_text": self.search_text,
            "bill_text": self.bill_text,
            "materials_json": [
                asdict(material) for material in self.materials
            ],
            "labor_cost": self.labor_cost,
            "material_cost": self.material_cost,
            "machine_cost": self.machine_cost,
            "management_fee": self.management_fee,
            "profit": self.profit,
            "measure_fee": self.measure_fee,
            "other_fee": self.other_fee,
            "tax": self.tax,
            "remarks": self.remarks,
            "tags": self.tags,
        }
        return record


@dataclass
class ParsedPricedBillDocument:
    file_path: str
    file_type: str
    project_name: str = ""
    specialty: str = ""
    items: list[ParsedPricedBillItem] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _classify_budget_row(serial: str, code: str, name: str, description: str) -> str:
    has_serial = bool(re.fullmatch(r"\d+", serial))
    has_bill_code = bool(re.fullmatch(r"\d{9,12}", code))
    has_desc = bool(description)
    cleaned_code = code.replace(" ", "")

    if name and ((has_bill_code and has_desc) or (has_serial and has_desc)):
        return "bill"

    if name and not description and not has_serial:
        if re.fullmatch(r"\d{7,8}", cleaned_code):
            return "quota"
        if _is_material_code(cleaned_code):
            return "material"
        if _is_quota_code(cleaned_code):
            return "quota"

    if cleaned_code and _is_quota_code(cleaned_code):
        return "quota"

    return "other"


def _iter_excel_bill_items(path: Path) -> tuple[list[ParsedPricedBillItem], list[str]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    items: list[ParsedPricedBillItem] = []
    warnings: list[str] = []
    try:
        for ws in wb.worksheets:
            current_bill: ParsedPricedBillItem | None = None
            current_section = ""
            saw_quota_like = False

            for row_index, row in enumerate(ws.iter_rows(min_row=1, values_only=True), start=1):
                if not row or all(cell is None for cell in row):
                    continue

                cells = list(row)
                if len(cells) < 9:
                    cells.extend([None] * (9 - len(cells)))

                serial = _clean_text(cells[0])
                code = _clean_text(cells[1])
                name = _clean_text(cells[2])
                description = _clean_text(cells[3])
                unit = _clean_text(cells[5])
                quantity = _coerce_float(cells[6])
                unit_price = _coerce_float(cells[7])

                if serial in {"序号", ""} and code in {"项目编码", "编码", ""} and name in {"项目名称", "名称", ""}:
                    continue

                row_type = _classify_budget_row(serial, code, name, description)
                if row_type == "bill":
                    if current_bill is not None:
                        items.append(current_bill)
                    bill_text = " ".join(part for part in [name, description] if part).strip()
                    current_bill = ParsedPricedBillItem(
                        boq_code=code,
                        boq_name_raw=name,
                        boq_name_normalized=_normalize_name(name),
                        feature_text=description,
                        unit=unit,
                        quantity=quantity,
                        composite_unit_price=unit_price,
                        source_sheet=ws.title,
                        source_row_no=row_index,
                        system_name=current_section,
                        bill_text=bill_text,
                        search_text=bill_text,
                    )
                    continue

                if row_type == "quota" and current_bill is not None:
                    saw_quota_like = True
                    current_bill.quotas.append(
                        ParsedQuota(
                            code=code,
                            name=name,
                            unit=unit,
                            quantity=quantity,
                        )
                    )
                    continue

                if row_type == "material" and current_bill is not None:
                    current_bill.materials.append(
                        ParsedMaterial(
                            code=code,
                            name=name,
                            unit=unit,
                            price=unit_price,
                        )
                    )
                    continue

                if name and len(name) >= 2 and not serial and not code and not description:
                    current_section = name

            if current_bill is not None:
                items.append(current_bill)

            if not saw_quota_like and items:
                warnings.append(f"sheet {ws.title} did not expose quota rows; kept bill-level prices only")
    finally:
        wb.close()

    return items, warnings


def _fallback_flat_excel(path: Path) -> tuple[list[ParsedPricedBillItem], list[str]]:
    reader = BillReader()
    rows = reader.read_file(str(path))
    items: list[ParsedPricedBillItem] = []
    for row in rows:
        name = _clean_text(row.get("name"))
        if not name:
            continue
        feature_text = _clean_text(row.get("description"))
        bill_text = " ".join(part for part in [name, feature_text] if part).strip()
        items.append(
            ParsedPricedBillItem(
                boq_code=_clean_text(row.get("code")),
                boq_name_raw=name,
                boq_name_normalized=_normalize_name(name),
                feature_text=feature_text,
                unit=_clean_text(row.get("unit")),
                quantity=row.get("quantity"),
                composite_unit_price=row.get("unit_price"),
                source_sheet=_clean_text(row.get("sheet_name")),
                source_row_no=row.get("source_row"),
                system_name=_clean_text(row.get("section")),
                bill_text=bill_text,
                search_text=row.get("search_text") or bill_text,
            )
        )
    return items, []


def _append_cost_fields(item: ParsedPricedBillItem, attrs: dict[str, str]) -> None:
    for field_name, keys in COST_ATTR_MAP.items():
        value = _find_numeric_attr(attrs, *keys)
        if value is not None:
            setattr(item, field_name, value)


def _append_zj_cost_fields(item: ParsedPricedBillItem, attrs: dict[str, str]) -> None:
    zj_map = {
        "labor_cost": ("\u4eba\u5de5\u8d39",),
        "material_cost": ("\u6750\u6599\u8d39",),
        "machine_cost": ("\u673a\u68b0\u8d39",),
        "management_fee": ("\u7ba1\u7406\u8d39",),
        "profit": ("\u5229\u6da6",),
        "other_fee": ("\u98ce\u9669\u8d39\u7528",),
        "tax": ("\u7a0e\u91d1",),
    }
    for field_name, keys in zj_map.items():
        value = _find_numeric_attr(attrs, *keys)
        if value is not None:
            setattr(item, field_name, value)


def _parse_zj_xml_items(root: ET.Element) -> tuple[list[ParsedPricedBillItem], str]:
    items: list[ParsedPricedBillItem] = []
    project_name = ""

    proj_info = root.find(ZJ_PROJECT_NAME_TAG)
    if proj_info is not None:
        project_name = _find_attr(proj_info.attrib, *ZJ_ATTR_KEYS["project_name"])

    for unit_elem in root.iter(ZJ_UNIT_TABLE_TAG):
        unit_name = _find_attr(unit_elem.attrib, *ZJ_ATTR_KEYS["unit_name"])
        for spec_elem in unit_elem.findall(ZJ_SPECIALTY_TAG):
            specialty_name = _find_attr(spec_elem.attrib, *ZJ_ATTR_KEYS["specialty_name"])
            specialty_type = _find_attr(spec_elem.attrib, *ZJ_ATTR_KEYS["specialty_type"])
            for table_elem in spec_elem:
                if table_elem.tag not in {
                    "\u5206\u90e8\u5206\u9879\u5de5\u7a0b\u91cf\u6e05\u5355\u8868",
                    "\u6280\u672f\u63aa\u65bd\u9879\u76ee\u6e05\u5355\u8868",
                }:
                    continue
                for title_elem in table_elem:
                    if not title_elem.tag.endswith("\u6807\u9898"):
                        continue
                    section_name = _find_attr(title_elem.attrib, *ZJ_ATTR_KEYS["section_name"])
                    for record in title_elem:
                        if record.tag not in ZJ_BILL_TAGS:
                            continue
                        item = ParsedPricedBillItem(
                            boq_code=_find_attr(record.attrib, *ZJ_ATTR_KEYS["boq_code"]),
                            boq_name_raw=_find_attr(record.attrib, *ZJ_ATTR_KEYS["boq_name"]),
                            boq_name_normalized=_normalize_name(_find_attr(record.attrib, *ZJ_ATTR_KEYS["boq_name"])),
                            feature_text=_find_attr(record.attrib, *ZJ_ATTR_KEYS["feature_text"]),
                            unit=_find_attr(record.attrib, *ZJ_ATTR_KEYS["unit"]),
                            quantity=_find_numeric_attr(record.attrib, *ZJ_ATTR_KEYS["quantity"]),
                            composite_unit_price=_find_numeric_attr(record.attrib, *ZJ_ATTR_KEYS["composite_unit_price"]),
                            specialty=specialty_type,
                            system_name=unit_name,
                            subsystem_name=specialty_name,
                            bill_text=" ".join(
                                part
                                for part in [
                                    _find_attr(record.attrib, *ZJ_ATTR_KEYS["boq_name"]),
                                    _find_attr(record.attrib, *ZJ_ATTR_KEYS["feature_text"]),
                                ]
                                if part
                            ).strip(),
                            remarks=section_name,
                        )
                        _append_zj_cost_fields(item, record.attrib)

                        for child in record:
                            if child.tag not in ZJ_ANALYSIS_TAGS:
                                continue
                            item.quotas.append(
                                ParsedQuota(
                                    code=_find_attr(child.attrib, *ZJ_ATTR_KEYS["quota_code"]),
                                    name=_find_attr(child.attrib, *ZJ_ATTR_KEYS["quota_name"]),
                                    unit=_find_attr(child.attrib, *ZJ_ATTR_KEYS["quota_unit"]),
                                    quantity=_find_numeric_attr(child.attrib, *ZJ_ATTR_KEYS["quota_quantity"]),
                                )
                            )
                            if item.composite_unit_price is None:
                                item.composite_unit_price = _find_numeric_attr(
                                    child.attrib, *ZJ_ATTR_KEYS["composite_unit_price"]
                                )
                            _append_zj_cost_fields(item, child.attrib)

                        if item.quotas:
                            item.quota_code = ",".join(q.code for q in item.quotas if q.code)
                            item.quota_name = " + ".join(q.name for q in item.quotas if q.name)
                        item.search_text = _extract_search_text(
                            item.boq_name_raw,
                            item.feature_text,
                            [q.code for q in item.quotas if q.code],
                            [q.name for q in item.quotas if q.name],
                        )
                        if item.boq_name_raw:
                            items.append(item)

    return items, project_name


def _load_xml_root(path: Path) -> tuple[object, list[str]]:
    warnings: list[str] = []
    try:
        tree = ET.parse(path)
        return tree.getroot(), warnings
    except ET.ParseError as exc:
        if LET is None:
            raise ET.ParseError(
                f"{exc}; install lxml to enable malformed XML recovery"
            ) from exc
        parser = LET.XMLParser(recover=True, huge_tree=True)
        tree = LET.parse(str(path), parser)
        warnings.append(f"xml recovered with lxml: {exc}")
        return tree.getroot(), warnings


def _parse_xml_items(path: Path) -> tuple[list[ParsedPricedBillItem], list[str]]:
    root, warnings = _load_xml_root(path)
    items: list[ParsedPricedBillItem] = []

    if root.tag == "\u6d59\u6c5f\u7701\u5efa\u8bbe\u5de5\u7a0b\u8ba1\u4ef7\u6210\u679c\u6587\u4ef6\u6570\u636e\u6807\u51c6":
        items, project_name = _parse_zj_xml_items(root)
        if not items:
            warnings.append(f"no bill items parsed from xml root <{root.tag}>")
        return items, warnings

    for elem in root.iter():
        tag = elem.tag.split("}")[-1]
        tag_upper = tag.upper()
        is_bill_tag = (
            tag_upper == "QDXM"
            or "清单项目" in tag
            or "清单项目表记录" in tag
            or "分部分项工程量清单" in tag
            or "措施子目清单" in tag
        )
        if not is_bill_tag:
            continue

        attrs = elem.attrib
        bill_name = _find_attr(attrs, "XMMC", "项目名称", "Mc", "名称")
        if not bill_name:
            continue

        feature_lines = []
        feature_attr = _find_attr(attrs, "XMTZ", "项目特征", "项目特征描述", "工作内容", "特征描述", "特征")
        if feature_attr:
            feature_lines.append(feature_attr)
        for feature_node in elem.findall(".//*"):
            child_tag = feature_node.tag.split("}")[-1]
            if "特征明细" in child_tag or "项目特征子目" in child_tag:
                content = _find_attr(feature_node.attrib, "内容", "content")
                if content:
                    feature_lines.append(content)

        item = ParsedPricedBillItem(
            boq_code=_find_attr(attrs, "XMBM", "项目编码", "Qdbm"),
            boq_name_raw=bill_name,
            boq_name_normalized=_normalize_name(bill_name),
            feature_text="\n".join(line for line in feature_lines if line),
            unit=_find_attr(attrs, "JLDW", "DW", "计量单位", "Dw"),
            quantity=_find_numeric_attr(attrs, "GCSL", "SL", "工程量", "数量", "用量", "DwQdSl"),
            composite_unit_price=_find_numeric_attr(attrs, *PRICE_ATTR_KEYS),
            bill_text=" ".join(part for part in [bill_name, " ".join(feature_lines)] if part).strip(),
        )
        _append_cost_fields(item, attrs)

        for quota_elem in elem.iter():
            quota_tag = quota_elem.tag.split("}")[-1]
            quota_tag_upper = quota_tag.upper()
            is_quota_tag = (
                quota_tag_upper in {"DEZM", "QDXDEZJMX"}
                or "定额子目" in quota_tag
                or "综合单价分析表" in quota_tag
            )
            if not is_quota_tag:
                continue
            quota_attrs = quota_elem.attrib
            quota_name = _find_attr(quota_attrs, "XMMC", "DEMC", "名称", "Mc")
            quota_code = _find_attr(quota_attrs, "DEBH", "DEBM", "定额编号", "编码", "Debm")
            if not quota_name and not quota_code:
                continue
            item.quotas.append(
                ParsedQuota(
                    code=quota_code,
                    name=quota_name,
                    unit=_find_attr(quota_attrs, "JLDW", "DW", "计量单位", "单位", "Dw"),
                    quantity=_find_numeric_attr(quota_attrs, "GCSL", "SL", "用量", "数量", "DwQdSl"),
                )
            )
            if item.composite_unit_price is None:
                item.composite_unit_price = _find_numeric_attr(quota_attrs, *PRICE_ATTR_KEYS)
            _append_cost_fields(item, quota_attrs)

        if item.quotas:
            item.quota_code = ",".join(q.code for q in item.quotas if q.code)
            item.quota_name = " + ".join(q.name for q in item.quotas if q.name)
        item.search_text = _extract_search_text(
            item.boq_name_raw,
            item.feature_text,
            [q.code for q in item.quotas if q.code],
            [q.name for q in item.quotas if q.name],
        )
        items.append(item)

    if not items:
        warnings.append(f"no bill items parsed from xml root <{root.tag}>")
    return items, warnings


def parse_priced_bill_document(
    file_path: str | Path,
    *,
    project_name: str = "",
    specialty: str = "",
) -> ParsedPricedBillDocument:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(path)

    ext = path.suffix.lower()
    if ext in {".xlsx", ".xlsm", ".xltx", ".xltm", ".xls"}:
        items, warnings = _iter_excel_bill_items(path)
        if not items:
            items, flat_warnings = _fallback_flat_excel(path)
            warnings.extend(flat_warnings)
        file_type = "excel"
    elif ext in {".xml", ".13jk"}:
        items, warnings = _parse_xml_items(path)
        file_type = "xml"
    else:
        raise ValueError(f"unsupported priced bill file: {path.name}")

    for item in items:
        if not item.bill_text:
            item.bill_text = " ".join(
                part for part in [item.boq_name_raw, item.feature_text] if part
            ).strip()
        if not item.search_text:
            item.search_text = _extract_search_text(
                item.boq_name_raw,
                item.feature_text,
                [q.code for q in item.quotas if q.code],
                [q.name for q in item.quotas if q.name],
            )
        if not item.boq_name_normalized:
            item.boq_name_normalized = _normalize_name(item.boq_name_raw)

    return ParsedPricedBillDocument(
        file_path=str(path),
        file_type=file_type,
        project_name=project_name or path.stem,
        specialty=specialty,
        items=items,
        warnings=warnings,
    )
