from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from src.source_pack import build_source_pack, make_source_id, read_text_with_fallbacks, safe_text, summarize_text, write_source_pack  # noqa: E402


def _extract_from_pdf(path: Path) -> str:
    import pdfplumber

    parts: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            text = safe_text(page.extract_text() or "")
            if text:
                parts.append(f"## Page {page_no}\n{text}")
    return "\n\n".join(parts).strip()


def _extract_from_docx(path: Path) -> str:
    import docx

    document = docx.Document(path)
    return "\n".join(paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()).strip()


def _extract_from_xlsx(path: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sections: list[str] = []
    try:
        for ws in wb.worksheets[:5]:
            rows: list[str] = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(cell).strip() for cell in row if str(cell or "").strip()]
                if cells:
                    rows.append("\t".join(cells))
            if rows:
                sections.append(f"# Sheet: {ws.title}\n" + "\n".join(rows[:200]))
    finally:
        wb.close()
    return "\n\n".join(sections).strip()


def _extract_from_csv(path: Path) -> str:
    rows: list[str] = []
    with path.open("r", encoding="utf-8", errors="ignore", newline="") as fh:
        reader = csv.reader(fh)
        for row in reader:
            cells = [cell.strip() for cell in row if cell.strip()]
            if cells:
                rows.append("\t".join(cells))
    return "\n".join(rows).strip()


def _extract_from_json(path: Path) -> str:
    return json.dumps(json.loads(read_text_with_fallbacks(path)), ensure_ascii=False, indent=2)


def extract_document_text(path: str | Path) -> str:
    source_path = Path(path)
    suffix = source_path.suffix.lower()
    if suffix in {".md", ".txt", ".text"}:
        return read_text_with_fallbacks(source_path)
    if suffix == ".json":
        return _extract_from_json(source_path)
    if suffix == ".csv":
        return _extract_from_csv(source_path)
    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        return _extract_from_xlsx(source_path)
    if suffix == ".docx":
        return _extract_from_docx(source_path)
    if suffix == ".pdf":
        return _extract_from_pdf(source_path)
    return read_text_with_fallbacks(source_path)


def ingest_document(*, input_path: str | Path, title: str = "", province: str = "", specialty: str = "", tags: list[str] | None = None, confidence: int = 85) -> dict[str, str]:
    source_path = Path(input_path).resolve()
    text = extract_document_text(source_path)
    resolved_title = safe_text(title) or source_path.stem
    source_id = make_source_id(source_kind="doc", title=resolved_title, source_path=source_path)
    pack = build_source_pack(
        source_id=source_id,
        source_kind="doc",
        title=resolved_title,
        summary=summarize_text(text),
        full_text_path="",
        evidence_refs=[str(source_path)],
        province=province,
        specialty=specialty,
        tags=(tags or []) + ["document"],
        confidence=confidence,
        metadata={"input_ext": source_path.suffix.lower()},
    )
    pack_path, text_path = write_source_pack(pack=pack, full_text=text)
    return {"source_id": source_id, "pack_path": str(pack_path), "full_text_path": str(text_path)}


def main() -> int:
    parser = argparse.ArgumentParser(description="Ingest one document into source pack format.")
    parser.add_argument("--input", required=True)
    parser.add_argument("--title", default="")
    parser.add_argument("--province", default="")
    parser.add_argument("--specialty", default="")
    parser.add_argument("--tags", default="")
    parser.add_argument("--confidence", type=int, default=85)
    args = parser.parse_args()

    result = ingest_document(
        input_path=args.input,
        title=args.title,
        province=args.province,
        specialty=args.specialty,
        tags=[item.strip() for item in args.tags.split(",") if item.strip()],
        confidence=args.confidence,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
